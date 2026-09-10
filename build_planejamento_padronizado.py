from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = "/run/media/heisenberg0x01/SSD Kingston/PROJECTS/PET/FORMACAO/gerador_codigo/Planejamento 01 padronizado.docx"


def shade(cell, color):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        pr.append(node)
    node.set(qn("w:fill"), color)


def cell_margins(cell, value=110):
    pr = cell._tc.get_or_add_tcPr()
    mar = pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        pr.append(mar)
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)


def table_borders(table):
    pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "8")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    pr.append(borders)


def bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    return p


def topic_header(doc, heading, date, time, member):
    p = doc.add_paragraph(heading, style="Heading 2")
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    values = (("Data", date), ("Horário", time), ("Integrante", member))
    widths = (Cm(4.0), Cm(5.0), Cm(8.0))
    for cell, (label, value), width in zip(table.rows[0].cells, values, widths):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, "F1F5F9")
        cell_margins(cell, 65)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(0)
        run = cp.add_run(f"{label}: ")
        run.bold = True
        run.font.color.rgb = RGBColor(36, 59, 90)
        cp.add_run(value)
        for run in cp.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(8.7)
    table_borders(table)
    return table


def topic(doc, heading, date, time, member, items, intro=None):
    topic_header(doc, heading, date, time, member)
    if intro:
        doc.add_paragraph(intro)
    for item in items:
        bullet(doc, item)


doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21.59), Cm(27.94)
section.top_margin, section.bottom_margin = Cm(1.5), Cm(1.5)
section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Aptos", Pt(9.8)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(3)
normal.paragraph_format.line_spacing = 1.0

title = doc.styles["Title"]
title.font.name, title.font.size, title.font.bold = "Aptos Display", Pt(25), True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(8)
ppr = title.element.get_or_add_pPr()
border = ppr.find(qn("w:pBdr"))
if border is not None:
    ppr.remove(border)

for name, size, before, after in (
    ("Heading 1", 17, 11, 6),
    ("Heading 2", 13, 7, 3),
    ("Heading 3", 11.5, 5, 2),
):
    style = doc.styles[name]
    style.font.name, style.font.size, style.font.bold = "Aptos Display", Pt(size), True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run("Planejamento da formação em Arduino ESP32 e impressão 3D")
r.font.name, r.font.size = "Aptos", Pt(8.5)
r.font.color.rgb = RGBColor(95, 95, 95)

doc.add_paragraph("Planejamento da formação em Arduino ESP32 e impressão 3D", style="Title")
p = doc.add_paragraph("Cronograma geral e conteúdos programáticos")
p.paragraph_format.space_after = Pt(12)
for r in p.runs:
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(75, 75, 75)
doc.add_paragraph(
    "O planejamento reúne as datas, os horários, os integrantes responsáveis e os conteúdos previstos para cada encontro da formação."
)

doc.add_heading("Cronograma geral", level=1)
headers = ["Horário", "10/09", "17/09", "24/09", "01/10", "Data a definir"]
data = [
    ["15h00", "Gabriel\nIntrodução e fundamentos físicos", "Ana Maria\nConteúdo a definir", "Gabriel\nHardware e ESP32", "A definir", "A definir"],
    ["15h45", "Gabriel\nIntrodução aos microcontroladores", "Ana Maria\nConteúdo a definir", "Rennê Bispo dos Santos\nFundamentos e prática inicial", "A definir", "A definir"],
    ["16h25", "Intervalo", "Intervalo", "Intervalo", "Intervalo", "Intervalo"],
    ["16h35", "Sérgio\nApresentação do Arduino", "Ana Maria\nConteúdo a definir", "Rennê Bispo dos Santos\nFuncionalidades do robô", "A definir", "Gabriel\nImpressora 3D"],
    ["17h15", "Sérgio\nPrática com Arduino", "Ana Maria\nConteúdo a definir", "Rennê Bispo dos Santos\nPrática no simulador", "A definir", "Gabriel\nImpressora 3D"],
    ["18h00", "Encerramento", "Encerramento", "Encerramento", "Encerramento", "Encerramento"],
]
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = True
for c, text in zip(table.rows[0].cells, headers):
    c.text = text
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(c, "243B5A")
    cell_margins(c)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in c.paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
for ridx, row in enumerate(data):
    cells = table.add_row().cells
    for cidx, (c, text) in enumerate(zip(cells, row)):
        c.text = text
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margins(c, 90)
        if ridx % 2:
            shade(c, "F1F5F9")
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cidx == 0 or text in ("Intervalo", "Encerramento", "A definir") else WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(7.5)
            if cidx == 0:
                run.bold = True
table_borders(table)

doc.add_paragraph("Os conteúdos ainda não informados permanecem identificados como a definir.")
doc.add_page_break()

doc.add_heading("Dia 1  10 de setembro de 2026", level=1)
p = doc.add_paragraph("Dos fundamentos físicos da eletricidade ao primeiro contato prático com o Arduino")
p.paragraph_format.space_after = Pt(7)
for r in p.runs:
    r.bold = True
    r.font.size = Pt(13)
doc.add_paragraph(
    "O encontro apresenta os conceitos de energia, eletricidade e controle necessários para compreender um microcontrolador e aplicá-los na simulação de circuitos com Arduino."
)

topic(doc, "1  Introdução e fundamentos físicos", "10/09/2026", "15h00 às 15h45", "Gabriel", [
    "Conceito de energia e princípio da conservação da energia",
    "Conversões de energia em motores, pilhas, LEDs e buzzers",
    "Manipulação e controle da energia com interruptores, resistores, transistores e potenciômetros",
    "Medição do ambiente com potenciômetros, sensores LDR e outros módulos",
    "Circuitos complexos, microcontroladores e computadores como formas sofisticadas de controle",
])

topic(doc, "2  Introdução aos microcontroladores", "10/09/2026", "15h45 às 16h25", "Gabriel", [
    "Saída digital com LEDs e resistores",
    "Entrada digital com interruptores e botões",
    "PWM para controle do brilho de LEDs e geração de sinais para buzzers",
    "Entrada analógica com potenciômetros e sensores LDR",
    "Módulos e atuadores: ponte H, sensor ultrassônico, sensor de presença e servomotor",
    "Identificação dos componentes, de suas funções, de seus terminais e das formas de conexão",
])

topic(doc, "3  Apresentação do Arduino", "10/09/2026", "16h35 às 17h15", "Sérgio", [
    "Funcionamento do Arduino",
    "Portas, conectores, entradas e saídas da placa",
    "Estrutura básica de um sketch e apresentação da Arduino IDE",
    "Montagem correta das conexões e envio de instruções por código",
])

topic(doc, "4  Prática no Tinkercad", "10/09/2026", "17h15 às 18h00", "Sérgio", [
    "LEDs", "Sensor ultrassônico", "Buzzer", "Fotoresistor LDR", "Microservomotor", "Display LCD",
], intro="Montagem e simulação de circuitos com os seguintes componentes:")

doc.add_page_break()
doc.add_heading("Dia 3  24 de setembro de 2026", level=1)
p = doc.add_paragraph("Fundamentos da robótica e prática de programação em blocos")
p.paragraph_format.space_after = Pt(7)
for r in p.runs:
    r.bold = True
    r.font.size = Pt(13)
doc.add_paragraph(
    "O encontro relaciona o funcionamento de um robô à programação em blocos e culmina na criação, troca e apresentação de rotinas no simulador."
)

topic(doc, "1  Fundamentos e prática inicial", "24/09/2026", "15h45 às 16h25", "Rennê Bispo dos Santos", [
    "Conceitos essenciais sobre o funcionamento de um robô",
    "Conexão entre o mundo real e o mundo virtual",
    "Base para compreender a programação em blocos",
])

topic(doc, "2  Funcionalidades do robô", "24/09/2026", "16h35 às 17h15", "Rennê Bispo dos Santos", [
    "Explicação teórica das funcionalidades do robô",
    "Aplicação das funcionalidades na programação",
    "Criação de rotinas básicas de funcionamento",
])

topic(doc, "3  Prática no simulador", "24/09/2026", "17h15 às 18h00", "Rennê Bispo dos Santos", [
    "Programação em blocos no simulador",
    "Criação de rotinas básicas por sorteio",
    "Troca de códigos entre os integrantes do PET",
    "Apresentação das diferentes soluções encontradas",
])

doc.core_properties.title = "Planejamento da formação em Arduino ESP32 e impressão 3D"
doc.core_properties.subject = "Cronograma geral e conteúdos programáticos"
doc.save(OUT)
print(OUT)
