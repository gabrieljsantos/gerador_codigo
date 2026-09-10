from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = "/run/media/heisenberg0x01/SSD Kingston/PROJECTS/PET/FORMACAO/gerador_codigo/Planejamento 01 revisado.docx"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def borders(table, color="D9D9D9", size="8"):
    tbl_pr = table._tbl.tblPr
    element = tbl_pr.find(qn("w:tblBorders"))
    if element is None:
        element = OxmlElement("w:tblBorders")
        tbl_pr.append(element)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)
        element.append(tag)


def margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.59)
sec.page_height = Cm(27.94)
sec.top_margin = Cm(2.0)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.0)
sec.right_margin = Cm(2.0)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.02

title = styles["Title"]
title.font.name = "Aptos Display"
title.font.size = Pt(26)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(8)
title_ppr = title.element.get_or_add_pPr()
title_borders = title_ppr.find(qn("w:pBdr"))
if title_borders is not None:
    title_ppr.remove(title_borders)

for style_name, size, before, after in (
    ("Heading 1", 18, 14, 7),
    ("Heading 2", 14, 9, 4),
    ("Heading 3", 12, 6, 2),
):
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Planejamento da formação em Arduino")
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)

p = doc.add_paragraph("Planejamento da formação em Arduino", style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
subtitle = doc.add_paragraph("Cronograma geral e conteúdo programático do primeiro encontro")
subtitle.paragraph_format.space_after = Pt(14)
run = subtitle.runs[0]
run.font.size = Pt(12)
run.font.color.rgb = RGBColor(70, 70, 70)

intro = doc.add_paragraph(
    "Este documento organiza o cronograma da formação e detalha o conteúdo do primeiro encontro, "
    "que parte dos fundamentos físicos da eletricidade e avança até a montagem de circuitos com Arduino no Tinkercad."
)
intro.paragraph_format.space_after = Pt(14)

doc.add_heading("Cronograma geral", level=1)
headers = ["Horário", "10/09", "17/09", "24/09", "01/10", "Data a definir"]
rows = [
    ["15h00", "Gabriel\nIntrodução e fundamentos físicos", "—", "Gabriel\nHardware", "—", "—"],
    ["15h45", "Gabriel\nIntrodução aos micro-\ncontroladores", "—", "Rennê", "—", "—"],
    ["16h25", "Intervalo", "Intervalo", "Intervalo", "Intervalo", "Intervalo"],
    ["16h35", "Sérgio\nApresentação do Arduino\nPrática com Arduino", "—", "Rennê", "—", "Gabriel\nModelagem 3D"],
    ["18h00", "Encerramento", "—", "—", "—", "—"],
]
table = doc.add_table(rows=1, cols=len(headers))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Cm(1.6), Cm(4.3), Cm(1.8), Cm(2.6), Cm(1.8), Cm(3.9)]
for i, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
    cell.width = width
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, "243B5A")
    margins(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
for ridx, values in enumerate(rows, start=1):
    cells = table.add_row().cells
    for i, (cell, text, width) in enumerate(zip(cells, values, widths)):
        cell.width = width
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        margins(cell, top=110, bottom=110)
        if ridx % 2 == 0:
            shade(cell, "F1F5F9")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or text in ("—", "Intervalo") else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.size = Pt(9)
            if i == 0:
                run.font.bold = True
borders(table)

note = doc.add_paragraph()
note.paragraph_format.space_before = Pt(7)
note.paragraph_format.space_after = Pt(0)
r = note.add_run("Observação: ")
r.bold = True
note.add_run("os campos sem atividade definida foram mantidos em aberto para preenchimento posterior.")

doc.add_page_break()

doc.add_heading("Dia 1  10 de setembro", level=1)
p = doc.add_paragraph("Dos fundamentos físicos da eletricidade ao primeiro contato prático com o Arduino")
p.paragraph_format.space_after = Pt(10)
for r in p.runs:
    r.font.size = Pt(13)
    r.font.bold = True

doc.add_paragraph(
    "O encontro apresenta os conceitos de energia, eletricidade e controle necessários para compreender o funcionamento de um microcontrolador. "
    "Ao final, os participantes aplicarão esses conceitos na simulação de circuitos com Arduino."
)

doc.add_heading("1  Introdução e fundamentos físicos", level=2)
doc.add_heading("Energia e suas conversões", level=3)
for text in (
    "Conceito de energia",
    "Princípio da conservação da energia",
    "Conversões de energia em motores, pilhas, LEDs e buzzers",
): add_bullet(doc, text)

doc.add_heading("Manipulação e controle da energia", level=3)
for text in ("Interruptores", "Corrente elétrica", "Resistores", "Transistores", "Potenciômetros"):
    add_bullet(doc, text)

doc.add_heading("Medição da natureza e do ambiente", level=3)
for text in ("Potenciômetros", "Sensores de luminosidade LDR", "Outros sensores e módulos"):
    add_bullet(doc, text)

doc.add_heading("Formas mais sofisticadas de controle", level=3)
for text in ("Circuitos complexos", "Microcontroladores", "Computadores"):
    add_bullet(doc, text)

doc.add_heading("2  Introdução aos microcontroladores", level=2)
sections = [
    ("Saída digital", ("Acionamento de LEDs", "Uso de resistores para limitar a corrente")),
    ("Entrada digital", ("Interruptores", "Botões")),
    ("PWM  Modulação por largura de pulso", ("Controle do brilho de LEDs", "Geração de sinais para buzzers")),
    ("Entrada analógica", ("Leitura de potenciômetros", "Leitura de sensores LDR")),
    ("Módulos e atuadores", ("Ponte H", "Sensor ultrassônico", "Sensor de presença", "Servomotor", "Outros módulos")),
    ("Identificação dos componentes", (
        "Reconhecimento de componentes, sensores e módulos",
        "Função e aplicação de cada componente",
        "Identificação dos terminais e das formas de conexão",
    )),
]
for heading, items in sections:
    doc.add_heading(heading, level=3)
    for item in items:
        add_bullet(doc, item)

doc.add_heading("3  Apresentação do Arduino", level=2)
for text in (
    "O que é o Arduino e como ele funciona",
    "Portas e conectores da placa",
    "Entradas e saídas digitais e analógicas",
    "Estrutura básica de um sketch",
    "Apresentação da Arduino IDE",
    "Montagem correta das conexões",
    "Envio de instruções por meio de código",
): add_bullet(doc, text)

doc.add_heading("4  Prática com Arduino no Tinkercad", level=2)
doc.add_paragraph("Montagem e simulação de circuitos utilizando os seguintes componentes:")
for text in ("LEDs", "Sensor ultrassônico", "Buzzer", "Fotoresistor LDR", "Microservomotor", "Display LCD"):
    add_bullet(doc, text)

doc.core_properties.title = "Planejamento da formação em Arduino"
doc.core_properties.subject = "Cronograma geral e conteúdo programático do primeiro encontro"
doc.save(OUT)
print(OUT)
